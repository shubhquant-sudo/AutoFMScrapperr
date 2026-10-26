"""Fetch history, write the workbook and keep a self-contained source package."""

import base64
import csv
import hashlib
import json
import os
import shutil
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from .data import BASE, acquire, normalize, ticker_name
from .model import make_model


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=True, allow_nan=False),
        encoding="utf-8",
    )


def write_csv(path, rows):
    def safe(v):
        return (
            "'" + v
            if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@"))
            else v
        )

    with Path(path).open("w", encoding="utf-8-sig", newline="") as file:
        csv.writer(file).writerows([[safe(v) for v in row] for row in rows])


def runtime():
    deps = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
    node = Path(os.environ.get("FM_NODE", str(deps / "node/bin/node.exe")))
    modules = Path(
        os.environ.get("FM_ARTIFACT_MODULES", str(deps / "node/node_modules"))
    )
    if not node.is_file() or not (modules / "@oai/artifact-tool").is_dir():
        raise ValueError(
            "The workbook writer runtime was not found. See README.md for FM_NODE and FM_ARTIFACT_MODULES."
        )
    return node, modules


def calculate(path):
    if os.name != "nt":
        return False, "Open the workbook in Microsoft Excel to calculate the formulas."
    shell = (
        Path(os.environ.get("SystemRoot", r"C:\Windows"))
        / "System32/WindowsPowerShell/v1.0/powershell.exe"
    )
    result = subprocess.run(
        [
            str(shell),
            "-NoProfile",
            "-NonInteractive",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(BASE / "autofmscrapper/recalculate.ps1"),
            "-Path",
            str(Path(path).resolve()),
        ],
        capture_output=True,
        check=False,
        text=True,
        timeout=120,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return result.returncode == 0, (
        result.stdout.strip() if result.returncode == 0 else result.stderr.strip()
    )[-1800:]


def author(spec, path, folder):
    if Path(path).exists():
        raise ValueError("This workbook already exists. Choose a new output folder.")
    spec_path = Path(folder) / "model_spec.json"
    write_json(spec_path, spec)
    node, modules = runtime()
    result = subprocess.run(
        [
            str(node),
            str(BASE / "autofmscrapper/write.mjs"),
            str(spec_path),
            str(Path(path).resolve()),
        ],
        env=dict(os.environ, FM_ARTIFACT_MODULES=str(modules)),
        capture_output=True,
        check=False,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=240,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    (Path(folder) / "authoring.log").write_text(
        result.stdout + "\n" + result.stderr, encoding="utf-8"
    )
    if result.returncode or not Path(path).is_file():
        raise ValueError(
            "Workbook creation failed. The source files and authoring log were retained in "
            + str(folder)
        )
    diagnostic = Path(str(path) + ".inspect.ndjson")
    if diagnostic.exists():
        shutil.move(str(diagnostic), str(Path(folder) / "authoring.inspect.ndjson"))


def inspect(path, spec, recalculated):
    """Check saved Excel values, source gaps and the blank-input contract."""
    cached_workbook = load_workbook(path, data_only=True)
    formula_workbook = load_workbook(path)
    try:

        def read_cell(target):
            sheet_name, cell_address = target.split("!", 1)
            return cached_workbook[sheet_name.strip("'")][cell_address].value

        formula_errors = [
            {"sheet": sheet.title, "cell": cell.coordinate, "error": cell.value}
            for sheet in cached_workbook
            for row in sheet
            for cell in row
            if cell.data_type == "e"
        ]
        inputs_blank = all(
            cached_workbook[item["sheet"]][item["cell"]].value is None
            for item in spec["inputs"].values()
        )
        checks = []
        for check in spec["checks"]:
            difference = read_cell(check["cell"])
            if difference is None:
                status = "Missing source / inputs"
            elif abs(difference) <= check["tolerance"]:
                status = "OK"
            else:
                status = "Review"
            checks.append({**check, "difference": difference, "status": status})

        last_actual_column = spec["last_actual_column"]
        statements = cached_workbook["Statements"]
        source_rows = {
            9: "Revenue",
            31: "Cash",
            34: "Net fixed assets",
            35: "Other assets / working-capital breakdown",
            36: "Total assets",
            39: "Debt",
            40: "Other liabilities / trade-payable breakdown",
            42: "Equity",
            43: "Total liabilities and equity",
            86: "Working-capital breakdown",
        }
        missing_source_fields = [
            label
            for row, label in source_rows.items()
            if not isinstance(
                statements.cell(row, last_actual_column).value, (int, float)
            )
        ]
        historical_columns = {chr(64 + column) for column in spec["historical_columns"]}
        historical_reviews = [
            check
            for check in checks
            if check["cell"].split("!")[1][0] in historical_columns
            and check["status"] == "Review"
        ]
        opening_notes = []
        for row, label in [
            (34, "Net fixed assets"),
            (35, "Other assets"),
            (39, "Debt"),
            (40, "Other liabilities"),
        ]:
            value = statements.cell(row, last_actual_column).value
            if isinstance(value, (int, float)) and value < 0:
                opening_notes.append(
                    label + " is negative; check the source classification."
                )

        ready = (
            recalculated
            and not formula_errors
            and inputs_blank
            and not formula_workbook._external_links
            and all(not sheet.sheet_view.showGridLines for sheet in formula_workbook)
        )
        return {
            "verified": bool(ready),
            "recalculated_by_excel": recalculated,
            "formula_errors": formula_errors,
            "all_forecast_and_dcf_inputs_blank": inputs_blank,
            "history_periods": spec["periods"],
            "missing_latest_source_fields": missing_source_fields,
            "historical_checks_to_review": historical_reviews,
            "opening_source_notes": opening_notes,
            "checks": checks,
            "outputs": {
                label: read_cell(target) for label, target in spec["outputs"].items()
            },
            "sheets": cached_workbook.sheetnames,
        }
    finally:
        cached_workbook.close()
        formula_workbook.close()


def _create_run_folders(ticker, output):
    """Give each run a new location so earlier work stays intact."""
    stamp = (
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ") + "_" + uuid4().hex[:6]
    )
    folder = Path(output).resolve() if output else BASE / "Output" / ticker / stamp
    package = folder.parent / (folder.name + ".zip")
    if folder.exists() or package.exists():
        raise ValueError(
            "The output folder or ZIP already exists. Choose a new folder so earlier work is preserved."
        )
    folder.mkdir(parents=True)
    sources = folder / "Sources"
    sources.mkdir()
    build = BASE / "cache/builds" / stamp
    build.mkdir(parents=True)
    return folder, sources, build, package


def _export_source_data(raw, snapshot, route, ticker, sources):
    """Save original bytes alongside normalized history and source references."""
    shutil.copyfile(snapshot, sources / "raw_source.json")
    digest = hashlib.sha256(Path(snapshot).read_bytes()).hexdigest()
    if raw.get("original_file_base64"):
        original = base64.b64decode(raw["original_file_base64"])
        if hashlib.sha256(original).hexdigest() != raw["input_sha256"]:
            raise ValueError("The imported file hash does not match the saved source.")
        (sources / ("Original_" + Path(raw["input_name"]).name)).write_bytes(original)
    data = normalize(raw, years=10000)
    data["notes"] = [
        s for s in data.get("notes", []) if not s.startswith("Requested 10000")
    ]
    write_json(sources / "Normalized_History.json", data)
    fields = sorted({k for row in data["history"] for k in row["values"]})
    write_csv(
        sources / "All_History.csv",
        [["period", *fields]]
        + [
            [row["period"], *[row["values"].get(k) for k in fields]]
            for row in data["history"]
        ],
    )
    names = ["period", "key", "field", "value", "status", "note"]
    write_csv(
        sources / "Field_Lineage.csv",
        [names] + [[row.get(k) for k in names] for row in data.get("lineage", [])],
    )
    write_json(
        sources / "Source_Manifest.json",
        {
            "ticker": ticker,
            "provider": raw["provider"],
            "source_url": raw.get("source_url"),
            "retrieved_utc": raw.get("retrieved_at_utc"),
            "basis": data["basis"],
            "route": route,
            "raw_file": "raw_source.json",
            "sha256": digest,
            "normalized_units": data["unit"],
            "shares_units": "millions",
        },
    )
    return data


def _write_run_archive(folder, package):
    """Include the workbook, instructions and source files in one archive."""
    (folder / "Read me.txt").write_text(
        (
            "Start on Assumptions. Check History, then enter the yellow forecast inputs "
            "and the DCF inputs. All forecast/valuation inputs are intentionally blank. "
            "Zero and missing are different.\n\nStatements contains the income statement, "
            "balance sheet and cash flow, with supporting calculations and checks below "
            "them. DCF uses that forecast. History contains all normalized annual "
            "observations; Raw statements retains the original source lines. Sources "
            "contains exact raw files, normalized history and the field/cell maps.\n\nIf a "
            "required historical value is missing, check the original filing and complete"
            " its yellow cell on History. Never replace a missing figure with zero unless"
            " it is actually zero. Source period labels and reporting basis are as "
            "supplied.\n\nThe forecast holds other assets/liabilities constant and excludes"
            " acquisitions, disposals, stock compensation, FX and additional noncash "
            "adjustments. It does not add debt automatically to hide negative cash.\n"
        ),
        encoding="utf-8",
    )
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as archive:
        for item in sorted(folder.rglob("*")):
            if item.is_file():
                archive.write(item, item.relative_to(folder).as_posix())


def generate(ticker, *, import_file=None, refresh=False, output=None, progress=None):
    ticker = ticker_name(ticker)
    notify = progress or (lambda step, text: None)
    folder, sources, build, package = _create_run_folders(ticker, output)
    notify(1, "Getting all available annual statements…")
    raw, snapshot, route = acquire(
        ticker,
        cache=BASE / "cache/snapshots",
        imports=BASE / "imports",
        screener=import_file,
        refresh=refresh,
    )
    data = _export_source_data(raw, snapshot, route, ticker, sources)
    notify(2, "Linking history and preparing empty forecast inputs…")
    spec = make_model(data, raw)
    path = folder / "AutoFMScrapper.xlsx"
    author(spec, path, build)
    notify(3, "Calculating and checking the workbook in Excel…")
    calculated, message = calculate(path)
    report = inspect(path, spec, calculated)
    report.update(
        ticker=ticker,
        company=data["company"],
        workbook=str(path),
        folder=str(folder),
        zip=str(package),
        created_utc=datetime.now(timezone.utc).isoformat(),
        source_route=route,
        message=message,
        unit=data["unit"],
    )
    write_json(folder / "Checks.json", report)
    write_json(
        sources / "Cell_Map.json",
        {
            "inputs": spec["inputs"],
            "history": spec["history_inputs"],
            "outputs": spec["outputs"],
        },
    )
    write_json(build / "report.json", report)
    if not report["verified"]:
        raise ValueError(
            "The workbook and sources were saved, but calculation verification needs attention. See "
            + str(folder / "Checks.json")
        )
    notify(4, "Saving the workbook and complete source package…")
    _write_run_archive(folder, package)
    report["zip"] = str(package)
    notify(5, "Ready. Open the workbook and enter your assumptions.")
    return report


def history():
    result = []
    for file in (BASE / "Output").glob("*/*/Checks.json"):
        try:
            report = json.loads(file.read_text(encoding="utf-8"))
            report.update(
                folder=str(file.parent.resolve()),
                workbook=str((file.parent / "AutoFMScrapper.xlsx").resolve()),
                zip=str((file.parent.parent / (file.parent.name + ".zip")).resolve()),
            )
            if report.get("verified") and Path(report["workbook"]).exists():
                result.append(report)
        except (ValueError, OSError):
            continue
    return sorted(result, key=lambda r: r["created_utc"], reverse=True)
